"""
Lambda function to generate Excel reports from Aurora database
Uploads the report to S3 for download
"""
import json
import os
import logging
import boto3
import psycopg
from datetime import datetime
from io import BytesIO
import sys

# Add lib directory to path for dependencies
sys.path.insert(0, '/var/task')

logger = logging.getLogger()
logger.setLevel(logging.INFO)

def get_secret(secret_arn, region):
    """Retrieve database credentials from Secrets Manager"""
    client = boto3.client('secretsmanager', region_name=region)
    response = client.get_secret_value(SecretId=secret_arn)
    return json.loads(response['SecretString'])

def fetch_resources_from_database(db_host, db_name, db_user, db_password, latest_only=True, account_ids=None):
    """Fetch resources from the database
    
    Args:
        db_host: Database host
        db_name: Database name
        db_user: Database user
        db_password: Database password
        latest_only: If True, only fetch resources from the latest discovery run
        account_ids: Optional list of account IDs to filter by
    """
    logger.info(f"Connecting to database: {db_host}")
    
    conn = psycopg.connect(
        host=db_host,
        port=5432,
        dbname=db_name,
        user=db_user,
        password=db_password
    )
    
    # Build account filter clause
    params = []
    account_filter = ""
    if account_ids:
        placeholders = ", ".join(["%s"] * len(account_ids))
        account_filter = f"AND r.account_id IN ({placeholders})"
        params = list(account_ids)
        logger.info(f"Filtering by {len(account_ids)} accounts: {account_ids}")

    if latest_only:
        # When filtering by accounts, scope the latest_date to those accounts too
        cte_filter = ""
        if account_ids:
            cte_placeholders = ", ".join(["%s"] * len(account_ids))
            cte_filter = f"WHERE account_id IN ({cte_placeholders})"
            # Double the params: first set for CTE, second set for main WHERE
            params = list(account_ids) + list(account_ids)

        query = f"""
            WITH latest_date AS (
                SELECT DATE(MAX(inserted_at)) as max_date
                FROM resources
                {cte_filter}
            )
            SELECT 
                r.resource_id,
                r.resource_type,
                r.resource_arn,
                r.region,
                r.account_id,
                r.name,
                r.tags,
                r.properties,
                r.discovered_at,
                r.last_seen_at,
                r.inserted_at
            FROM resources r, latest_date
            WHERE DATE(r.inserted_at) = latest_date.max_date
            {account_filter}
            ORDER BY r.account_id, r.region, r.resource_type, r.resource_id
        """
    else:
        where_clause = f"WHERE r.account_id IN ({', '.join(['%s'] * len(account_ids))})" if account_ids else ""
        query = f"""
            SELECT 
                r.resource_id,
                r.resource_type,
                r.resource_arn,
                r.region,
                r.account_id,
                r.name,
                r.tags,
                r.properties,
                r.discovered_at,
                r.last_seen_at,
                r.inserted_at
            FROM resources r
            {where_clause}
            ORDER BY r.account_id, r.region, r.resource_type, r.resource_id
        """
    
    resources = []
    with conn.cursor() as cur:
        cur.execute(query, params if params else None)
        rows = cur.fetchall()
        
        for row in rows:
            resource_dict = {
                'arn': row[2] if row[2] else '',
                'resource_type': row[1],
                'region': row[3],
                'account_id': row[4],
                'name': row[5],
                'tags': row[6] if row[6] else {},
                'configuration': row[7] if row[7] else {},
                'resource_id': row[0],
                'discovered_at': row[8].isoformat() if row[8] else None,
                'last_seen_at': row[9].isoformat() if row[9] else None,
                'inserted_at': row[10].isoformat() if row[10] else None,
            }
            resources.append(resource_dict)
    
    conn.close()
    logger.info(f"Fetched {len(resources)} resources from database (latest_only={latest_only})")
    return resources

def generate_excel_report(resources):
    """Generate Excel report from resources"""
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.utils.dataframe import dataframe_to_rows
    
    logger.info("Generating Excel report...")
    
    # Create DataFrame
    df = pd.DataFrame(resources)
    
    if df.empty:
        logger.warning("No resources to report")
        return None
    
    # Create Excel file in memory
    output = BytesIO()
    
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        # 1. Executive Summary
        # Get discovery timestamp from first resource
        discovery_timestamp = 'N/A'
        if 'inserted_at' in df.columns and not df['inserted_at'].isna().all():
            discovery_timestamp = df['inserted_at'].iloc[0] if len(df) > 0 else 'N/A'
        
        summary_data = {
            'Metric': [
                'Total Resources',
                'Unique Resource Types',
                'Accounts',
                'Regions',
                'Discovery Run',
                'Report Generated'
            ],
            'Value': [
                len(df),
                df['resource_type'].nunique(),
                df['account_id'].nunique(),
                df['region'].nunique(),
                discovery_timestamp,
                datetime.now().strftime('%Y-%m-%d %H:%M:%S UTC')
            ]
        }
        summary_df = pd.DataFrame(summary_data)
        summary_df.to_excel(writer, sheet_name='Executive Summary', index=False)
        
        # 2. All Resources
        # Select key columns for the main view
        main_cols = ['account_id', 'region', 'resource_type', 'name', 'arn', 'resource_id', 'inserted_at']
        available_cols = [col for col in main_cols if col in df.columns]
        df[available_cols].to_excel(writer, sheet_name='All Resources', index=False)
        
        # 3. By Resource Type
        type_summary = df.groupby('resource_type').size().reset_index(name='count')
        type_summary = type_summary.sort_values('count', ascending=False)
        type_summary.to_excel(writer, sheet_name='By Type', index=False)
        
        # 4. By Account
        account_summary = df.groupby('account_id').agg({
            'resource_id': 'count',
            'resource_type': 'nunique',
            'region': 'nunique'
        }).reset_index()
        account_summary.columns = ['account_id', 'total_resources', 'unique_types', 'regions']
        account_summary.to_excel(writer, sheet_name='By Account', index=False)
        
        # 5. By Region
        region_summary = df.groupby('region').size().reset_index(name='count')
        region_summary = region_summary.sort_values('count', ascending=False)
        region_summary.to_excel(writer, sheet_name='By Region', index=False)
    
    output.seek(0)
    logger.info("Excel report generated successfully")
    return output.getvalue()

def upload_to_s3(file_content, bucket_name, file_key):
    """Upload file to S3"""
    from botocore.config import Config
    
    # Configure S3 client with signature version 4
    s3_config = Config(signature_version='s3v4')
    s3 = boto3.client('s3', config=s3_config)
    
    s3.put_object(
        Bucket=bucket_name,
        Key=file_key,
        Body=file_content,
        ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    
    # Generate presigned URL (valid for 15 minutes to avoid clock skew issues)
    # Include Content-Disposition to force browser download with proper filename
    download_name = file_key.split('/')[-1]  # e.g. "CloudAuditor_Report_20260218_223652.xlsx"
    url = s3.generate_presigned_url(
        'get_object',
        Params={
            'Bucket': bucket_name,
            'Key': file_key,
            'ResponseContentDisposition': f'attachment; filename="{download_name}"',
        },
        ExpiresIn=900  # 15 minutes
    )
    
    return url

def lambda_handler(event, context):
    """
    Generate Excel report from Aurora database and upload to S3
    """
    logger.info(f"Event: {json.dumps(event)}")
    
    try:
        # Get configuration from environment
        secret_arn = os.environ['DB_SECRET_ARN']
        db_host = os.environ['DB_HOST']
        db_name = os.environ['DB_NAME']
        region = os.environ['AWS_REGION']
        bucket_name = os.environ.get('REPORT_BUCKET', 'cloudauditor-reports')
        
        # Get database credentials
        secret = get_secret(secret_arn, region)
        db_user = secret['username']
        db_password = secret['password']
        
        # Fetch resources from database (optionally scoped to specific accounts)
        account_ids = event.get('account_ids')
        resources = fetch_resources_from_database(db_host, db_name, db_user, db_password,
                                                  account_ids=account_ids)
        
        if not resources:
            return {
                'statusCode': 200,
                'body': json.dumps({
                    'success': True,
                    'message': 'No resources found in database',
                    'resource_count': 0
                })
            }
        
        # Generate Excel report
        excel_content = generate_excel_report(resources)
        
        if not excel_content:
            return {
                'statusCode': 500,
                'body': json.dumps({
                    'success': False,
                    'error': 'Failed to generate Excel report'
                })
            }
        
        # Upload to S3
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        file_key = f"reports/CloudAuditor_Report_{timestamp}.xlsx"
        
        download_url = upload_to_s3(excel_content, bucket_name, file_key)
        
        logger.info(f"Report uploaded successfully: {file_key}")
        
        return {
            'statusCode': 200,
            'body': json.dumps({
                'success': True,
                'message': f'Excel report generated with {len(resources)} resources',
                'resource_count': len(resources),
                'download_url': download_url,
                's3_bucket': bucket_name,
                's3_key': file_key,
                'expires_in_seconds': 900  # 15 minutes
            })
        }
        
    except Exception as e:
        logger.error(f"Error generating report: {e}")
        import traceback
        traceback.print_exc()
        
        return {
            'statusCode': 500,
            'body': json.dumps({
                'success': False,
                'error': str(e)
            })
        }
